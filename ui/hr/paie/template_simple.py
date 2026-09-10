"""موديل «Bulletin simple (CNAS)» — تخطيط + معاينة + إخراج Word/PDF.

التخطيط والترتيب على غرار نموذج كشف CNAS القصير المرجعي (رأس المكتب،
شريط العنوان الأسود، صندوق هوية الأجير، جدول
CODE/LIBELLÉ/N-BASE/TAUX/GAIN/RETENUE، سطر TOTAL، خانة NET À PAYER
السوداء). لا يُنقل أي محتوى من النموذج المرجعي — القياسات والترتيب فقط،
والبيانات كلها من الشاشة.
"""
import os
import tkinter.font as tkfont

from programme.payroll import registry
from programme.payroll.calc import fmt_montant
from ui.hr.constants import MOIS_FR, PAIE_DEFAULT_CODES
from ui.hr.render import TemplateNotReady

# خط الاستمارة الموحّد — نفس عائلة الخط للجزء الثابت المرسوم وخانات
# الإدخال فوقه (لا Courier) حتى تطابق كتابة المستخدم مظهر النموذج الأصلي.
FORM_FONT = "Helvetica"

# ------- تخطيط الصفحة (مليمترات على ورقة 210×297) -------
MARGIN_L = 14.0
MARGIN_R = 14.0
MARGIN_T = 12.0
CONTENT_W = 210.0 - MARGIN_L - MARGIN_R          # 182

# حدود الأعمدة نسبةً إلى عرض المحتوى (0..1)
COLS = [
    ("code",    0.000, 0.077, "c", "CODE"),
    ("libelle", 0.077, 0.450, "l", "LIBELLÉ"),
    ("nbase",   0.450, 0.615, "r", "N/BASE"),
    ("taux",    0.615, 0.755, "r", "TAUX"),
    ("gain",    0.755, 0.878, "r", "GAIN"),
    ("retenue", 0.878, 1.000, "r", "RETENUE"),
]

BAND_TITLE_Y = 40.0
BAND_TITLE_H = 8.0
IDENT_Y = 50.0
IDENT_H = 52.0
TABLE_HEAD_Y = 105.0
ROW_H = 6.2
TABLE_BOTTOM = 250.0

COMPUTED_COLOR = "#1a56b0"   # لون القيَم المحسوبة تلقائياً (تمييزها عن المكتوب يدوياً)

# ------- شبكة التحرير الثابتة (نمط CD: حقول فوق الاستمارة مباشرة) -------
# جسم الجدول = 10 أسطر ثابتة الترتيب (لا تتحرك) — كل موضع معروف مسبقاً:
BODY_TOP = TABLE_HEAD_Y + ROW_H + 1.0
N_PRIME_SLOTS = 3
N_AUTRE_SLOTS = 2
ROW_SALAIRE = 0
ROW_PRIME_0 = 1                      # 1..3
ROW_CNAS = 4
ROW_IRG = 5
ROW_PANIER = 6
ROW_TRANSPORT = 7
ROW_AUTRE_0 = 8                      # 8..9
N_BODY_ROWS = 10
FORM_TABLE_BOTTOM = BODY_TOP + N_BODY_ROWS * ROW_H
FORM_TOTAL_Y = FORM_TABLE_BOTTOM + 1.0
FORM_NET_Y = FORM_TOTAL_Y + ROW_H + 4.0


def _colf(col_key):
    for key, f0, f1, al, _t in COLS:
        if key == col_key:
            return f0, f1, al
    raise KeyError(col_key)


def _cell_mm(col_key, row_idx, pad=2.0):
    """(x_mm من حافة المحتوى اليسرى، y_mm من أعلى الورقة، w_mm) لخانة جدول."""
    f0, f1, _al = _colf(col_key)
    x = f0 * CONTENT_W + pad
    w = (f1 - f0) * CONTENT_W - 2 * pad
    y = BODY_TOP + row_idx * ROW_H + 0.7
    return x, y, w


class Slot:
    __slots__ = ("key", "x_mm", "y_mm", "w_mm", "align", "maxlen", "kind",
                 "font_mm", "bold", "h_mm", "on_band", "baseline_mm", "fit_maxlen")

    def __init__(self, key, x_mm, y_mm, w_mm, align="l", maxlen=40, kind="text",
                 font_mm=3.2, bold=False, h_mm=None, on_band=False, baseline_mm=None,
                 fit_maxlen=False):
        self.key, self.x_mm, self.y_mm, self.w_mm = key, x_mm, y_mm, w_mm
        self.align, self.maxlen, self.kind = align, maxlen, kind
        self.font_mm, self.bold, self.h_mm = font_mm, bold, h_mm
        # fit_maxlen: يُضبط عرض الإطار وقت التخطيط إلى قدر maxlen حرفاً
        # بالخط الحالي بالضبط (لا أكثر) — قاعدة "حدود الحقل = عدد الأحرف
        # المسموح" (نفس cd-field-maxlen-principle).
        self.fit_maxlen = fit_maxlen
        # on_band: خانة فوق الشريط الأسود — خلفية سوداء وكتابة بيضاء
        self.on_band = on_band
        # baseline_mm: لو مُحدَّد، تُوضَع الخانة بحيث يقع خط أساس نصّها
        # على هذي القيمة بالضبط (يُحسب من font ascent، لا تخمين) — نفس خط
        # الأساس تُرسَم عليه التسمية الثابتة → محاذاة حقيقية على كل زوم.
        self.baseline_mm = baseline_mm


# ---- مستطيل هوية الأجير: 5 أسطر، عمودان + سطر ميلاد كامل العرض ----
# _VAL_L موحّد لكل حقول العمود الأيسر (تبدأ كلها تحت بعضها على نفس x)،
# ومختار ليسع أطول تسمية «DATE DE NAISSANCE :».
_LBL_L, _VAL_L, _W_L = 4.0, 38.0, 49.0
_LBL_R, _VAL_R, _W_R = 93.0, 124.0, 55.0


def _ident_row_y(r):
    return IDENT_Y + 9.0 + r * 9.0


# (key, label, x_label, x_value, w_value, row, maxlen, kind) — label=None: خانة تتبع سابقتها
# ملاحظات ترتيب:
#  - العمود الأيسر (بداية خاناته موحّدة عند _VAL_L): NOM ← DATE DE
#    NAISSANCE ← N° SS ← DATE D'ENTRÉE ← FONCTION — كلها تبدأ على نفس
#    خط البداية، تحت بعضها (بطلب صريح).
#  - العمود الأيمن (عند _VAL_R): PRÉNOM ← SIT. FAMILIALE ← MATRICULE.
#  - SIT. FAMILIALE: حرف واحد فقط (C/D/V/M) — خانة ضيّقة + قائمة منبثقة.
#  - N° SS: 12 رقماً بتجميع 2‑4‑4‑2 (+ مفتاح «/XX» اختياري) — نفس منطق N° ADHÉRENT.
#  - السطر 4: FONCTION وحده بعرض كامل (المسمّيات الوظيفية تطول).
IDENT_FIELDS = [
    ("nom",                 "NOM",              _LBL_L, _VAL_L, _W_L,  0, 24, "upper_alpha"),
    ("prenom",              "PRÉNOM",           _LBL_R, _VAL_R, _W_R,  0, 24, "titlecase"),
    ("date_naissance",      "DATE DE NAISSANCE", _LBL_L, _VAL_L, 20.0, 1, 10, "date_masked"),
    ("lieu_naissance",      None,               None,   64.0,   59.0, 1, 40, "upper_alpha"),
    ("num_ss",              "N° SS",            _LBL_L, _VAL_L, _W_R,  2, 13, "num_ss"),
    ("situation_familiale", "SIT. FAMILIALE",   _LBL_R, _VAL_R, 14.0, 2, 1,  "famille"),
    ("date_embauche",       "DATE D'ENTRÉE",    _LBL_L, _VAL_L, 20.0, 3, 10, "date_masked"),
    ("matricule",           "MATRICULE",        _LBL_R, _VAL_R, _W_L,  3, 20, "text"),
    ("fonction",            "FONCTION",         _LBL_L, _VAL_L, 84.0,  4, 40, "text"),
]

# حقول هوية إطارها يُضبط إلى قدر محتواها الأقصى بالضبط وقت التخطيط
# (لا عرض ثابت مبالَغ فيه بفراغ ذيلي) — راجع Slot.fit_maxlen و_relayout.
# للحقول المجمَّعة (num_ss) يُقاس أقصى صيغة منسَّقة لا maxlen الخام.
_FIT_MAXLEN_KEYS = {"matricule", "fonction", "num_ss"}
IDENT_FIELD_H = 5.0          # ارتفاع خانة الهوية الاسمي (مم) — مرجعي فقط؛
# الخانات ذات baseline_mm تأخذ ارتفاعها الطبيعي وقت التخطيط (نفس الخط ⇒
# نفس الارتفاع تلقائياً)، فلا يُفرَض هذا الرقم كي لا يُوسَّط النص عمودياً
# ويهبط تحت خط الأساس (راجع _relayout بـbulletin_paie.py).

# خطوط الأساس (مم من أعلى الورقة) لحقول الترويسة — التسمية الثابتة تُرسَم
# على نفس القيمة، والخانة تُوضَع بحيث يقع نصّها عليها بالضبط.
RAISON_BASELINE = 10.5
ADRESSE_BASELINE = 18.5
ADHERENT_BASELINE = 26.5
BAND_BASELINE = BAND_TITLE_Y + BAND_TITLE_H / 2.0 + 1.6


def build_field_slots():
    """كل خانات الإدخال فوق الاستمارة — بالمليمتر في نفس إطار paint_form.

    الترتيب هنا = ترتيب التنقّل بـTab/Enter (قراءة الاستمارة من الأعلى).
    """
    slots = [
        # رأس المكتب — من الهامش للهامش (أسماء/عناوين طويلة)، بمقاسات الأصل
        Slot("emp_raison_sociale", 0.0, 0.0, CONTENT_W, "l", 46, "upper",
             font_mm=5.0, bold=True, h_mm=7.6, baseline_mm=RAISON_BASELINE),
        Slot("emp_adresse", 0.0, 0.0, CONTENT_W, "l", 95, "text",
             font_mm=3.4, bold=True, h_mm=5.2, baseline_mm=ADRESSE_BASELINE),
        Slot("emp_cnas", 24.0, 0.0, 62.0, "l", 16, "adherent",
             font_mm=3.4, bold=True, h_mm=5.2, baseline_mm=ADHERENT_BASELINE,
             fit_maxlen=True),
        Slot("mois", CONTENT_W * 0.52, 0.0, 34.0, "l", 12, "month",
             font_mm=5.0, bold=True, h_mm=7.0, on_band=True, baseline_mm=BAND_BASELINE),
        Slot("annee", CONTENT_W * 0.52 + 35.0, 0.0, 18.0, "l", 4, "year",
             font_mm=5.0, bold=True, h_mm=7.0, on_band=True, baseline_mm=BAND_BASELINE),
    ]
    for key, _lbl, _xl, xv, wv, row, mlen, kind in IDENT_FIELDS:
        slots.append(Slot(f"id_{key}", xv, 0.0, wv, "l", mlen, kind,
                          font_mm=3.2, bold=True, h_mm=IDENT_FIELD_H,
                          baseline_mm=_ident_row_y(row),
                          fit_maxlen=key in _FIT_MAXLEN_KEYS))
    # سطر SALAIRE DE BASE
    x, y, w = _cell_mm("nbase", ROW_SALAIRE)
    slots.append(Slot("jours", x, y, w, "r", 6, "num"))
    x, y, w = _cell_mm("gain", ROW_SALAIRE)
    slots.append(Slot("salaire_base", x, y, w, "r", 12, "num"))
    # أسطر المنح
    for k in range(N_PRIME_SLOTS):
        r = ROW_PRIME_0 + k
        x, y, w = _cell_mm("code", r)
        slots.append(Slot(f"prime{k}_code", x, y, w, "c", 5))
        x, y, w = _cell_mm("libelle", r)
        slots.append(Slot(f"prime{k}_lib", x, y, w, "l", 30))
        x, y, w = _cell_mm("gain", r)
        slots.append(Slot(f"prime{k}_montant", x, y, w, "r", 12, "num"))
    # panier / transport
    x, y, w = _cell_mm("gain", ROW_PANIER)
    slots.append(Slot("panier", x, y, w, "r", 12, "num"))
    x, y, w = _cell_mm("gain", ROW_TRANSPORT)
    slots.append(Slot("transport", x, y, w, "r", 12, "num"))
    # اقتطاعات أخرى
    for k in range(N_AUTRE_SLOTS):
        r = ROW_AUTRE_0 + k
        x, y, w = _cell_mm("code", r)
        slots.append(Slot(f"autre{k}_code", x, y, w, "c", 5))
        x, y, w = _cell_mm("libelle", r)
        slots.append(Slot(f"autre{k}_lib", x, y, w, "l", 30))
        x, y, w = _cell_mm("retenue", r)
        slots.append(Slot(f"autre{k}_montant", x, y, w, "r", 12, "num"))
    return slots


FIELD_SLOTS = build_field_slots()


def row1_layout(canvas, scale):
    """موضعا «à» وحقل المكان في سطر تاريخ الميلاد — يُحسبان من قياس الخط
    الفعلي (لا ثوابت): نهاية خانة التاريخ + مسافتان → «à» → مسافتان →
    حقل المكان. نفس الدالة يستدعيها paint_form (لرسم «à») و_relayout
    (لموضع خانة المكان) فيتطابقان."""
    fs = max(int(3.2 * scale * 0.62), 6)
    f = tkfont.Font(root=canvas, font=(FORM_FONT, fs, "bold"))
    date_w_mm = f.measure("00/00/0000") / scale + 1.8   # + هامش الخانة
    sp = f.measure(" ") / scale
    a_x = _VAL_L + date_w_mm + 2 * sp
    lieu_x = a_x + f.measure("à") / scale + 2 * sp
    return a_x, lieu_x


def slot_px(page_box, scale, slot):
    """(x, y, w, h) بالبكسل على اللوحة لخانة إدخال."""
    x0, y0 = page_box[0], page_box[1]
    h_mm = slot.h_mm if slot.h_mm is not None else (ROW_H - 0.8)
    return (
        x0 + (MARGIN_L + slot.x_mm) * scale,
        y0 + slot.y_mm * scale,
        slot.w_mm * scale,
        h_mm * scale,
    )


def paint_form(canvas, page_box, scale, pin, res, employer, employee):
    """يرسم الجزء الثابت + القيَم المحسوبة فقط (بلا القيَم المكتوبة يدوياً —
    مكانها خانات إدخال حيّة فوق اللوحة)."""
    x0, y0, _x1, _y1 = page_box

    def X(mm):
        return x0 + (MARGIN_L + mm) * scale

    def Y(mm):
        return y0 + mm * scale

    def F(mm_h, bold=False):
        size = max(int(mm_h * scale * 0.62), 6)
        return (FORM_FONT, size, "bold") if bold else (FORM_FONT, size)

    def col_x(frac):
        return x0 + (MARGIN_L + frac * CONTENT_W) * scale

    def base_text(mm_x, baseline_mm, text, font_tuple, fill="#202124"):
        """يرسم نصّاً ثابتاً بحيث يقع خط أساسه على baseline_mm بالضبط —
        نفس خط الأساس تُوضَع عليه الخانة (محاذاة حقيقية، لا تخمين)."""
        desc = tkfont.Font(root=canvas, font=font_tuple).metrics("descent")
        canvas.create_text(X(mm_x), Y(baseline_mm) + desc, text=text,
                           anchor="sw", font=font_tuple, fill=fill)

    ink = "#202124"
    cc = COMPUTED_COLOR

    # رأس المكتب — التسمية ثابتة على خط أساس حقل الانتساب، بنفس مقاس
    # خطّ خانته (3.4mm) حتى تبدو التسمية والخانة على سطر واحد فعلاً.
    base_text(0, ADHERENT_BASELINE, "N° ADHÉRENT", F(3.4, bold=True))

    # شريط العنوان الأسود
    canvas.create_rectangle(X(0), Y(BAND_TITLE_Y), X(CONTENT_W), Y(BAND_TITLE_Y + BAND_TITLE_H),
                            fill="#111111", outline="")
    base_text(3, BAND_BASELINE, "BULLETIN DE PAIE", F(5.0, bold=True), fill="white")

    # صندوق الهوية — التسميات فقط، كلها على خط أساس سطرها (تطابق الخانات)
    canvas.create_rectangle(X(0), Y(IDENT_Y), X(CONTENT_W), Y(IDENT_Y + IDENT_H), outline=ink)
    # التسميات بنفس مقاس خطّ خاناتها (3.2mm) — لا أصغر — حتى تقرأ التسمية
    # والخانة كسطر واحد متّسق، لا سطرين بمقاسين.
    for _key, lbl, xl, xv, wv, row, _ml, kind in IDENT_FIELDS:
        if lbl:
            base_text(xl, _ident_row_y(row), f"{lbl} :", F(3.2, bold=True))
        if kind == "famille":
            # سهم قائمة صغير مباشرةً بعد الخانة (خارجها، وإلا اختفى تحت
            # ودجت الإدخال المعتِم فوقه) يدلّ أنها قابلة للاختيار من قائمة
            base_text(xv + wv + 1.2, _ident_row_y(row), "▾", F(3.0), fill="#5f6368")
    a_x, _lieu_x = row1_layout(canvas, scale)
    base_text(a_x, _ident_row_y(1), "à", F(3.2, bold=True))

    # ترويسة الجدول + الأعمدة
    hy0, hy1 = Y(TABLE_HEAD_Y), Y(TABLE_HEAD_Y + ROW_H + 1)
    canvas.create_rectangle(X(0), hy0, X(CONTENT_W), hy1, outline=ink)
    for _key, f0, _f1, _al, title in COLS:
        canvas.create_line(col_x(f0), hy0, col_x(f0), Y(FORM_TABLE_BOTTOM), fill=ink)
        canvas.create_text(col_x(f0) + 2.5 * scale, (hy0 + hy1) / 2, text=title,
                           anchor="w", font=F(3.0, bold=True), fill=ink)
    canvas.create_line(X(CONTENT_W), hy0, X(CONTENT_W), Y(FORM_TABLE_BOTTOM), fill=ink)

    # خطوط الأسطر الأفقية
    for r in range(N_BODY_ROWS + 1):
        yy = Y(BODY_TOP + r * ROW_H)
        canvas.create_line(X(0), yy, X(CONTENT_W), yy, fill="#c8c8c8")

    def row_mid(r):
        return Y(BODY_TOP + r * ROW_H + ROW_H / 2)

    def cell_right(col_key, r, text, color=ink, bold=False):
        _f0, f1, _al = _colf(col_key)
        canvas.create_text(col_x(f1) - 2.5 * scale, row_mid(r), text=text, anchor="e",
                           font=F(3.2, bold=bold), fill=color)

    def cell_left(col_key, r, text, color=ink, bold=False):
        f0, _f1, _al = _colf(col_key)
        canvas.create_text(col_x(f0) + 2.5 * scale, row_mid(r), text=text, anchor="w",
                           font=F(3.2, bold=bold), fill=color)

    def cell_center(col_key, r, text, color=ink, bold=False):
        f0, f1, _al = _colf(col_key)
        canvas.create_text((col_x(f0) + col_x(f1)) / 2, row_mid(r), text=text, anchor="center",
                           font=F(3.2, bold=bold), fill=color)

    c = PAIE_DEFAULT_CODES
    # الأسطر ذات التسمية الثابتة
    cell_center("code", ROW_SALAIRE, c["salaire_base"])
    cell_left("libelle", ROW_SALAIRE, "SALAIRE DE BASE")
    cell_center("code", ROW_CNAS, c["cnas"])
    cell_left("libelle", ROW_CNAS, "RETENUE SÉCU. SOCIALE")
    cell_center("code", ROW_IRG, c["irg"])
    cell_left("libelle", ROW_IRG, "RETENUE IRG")
    cell_center("code", ROW_PANIER, c["panier"])
    cell_left("libelle", ROW_PANIER, "PANIER")
    cell_center("code", ROW_TRANSPORT, c["transport"])
    cell_left("libelle", ROW_TRANSPORT, "(R+) TRANSPORT")

    # القيَم المحسوبة تلقائياً (بلون مميّز)
    if res is not None:
        cell_right("nbase", ROW_CNAS, fmt_montant(res.base_cnas), cc)
        cell_right("taux", ROW_CNAS, "9,00", cc)
        cell_right("retenue", ROW_CNAS, fmt_montant(res.retenue_cnas), cc)
        cell_right("nbase", ROW_IRG, fmt_montant(res.base_irg), cc)
        cell_right("retenue", ROW_IRG, fmt_montant(res.retenue_irg), cc)

        # سطر TOTAL
        ty0 = Y(FORM_TOTAL_Y)
        canvas.create_line(X(0), ty0, X(CONTENT_W), ty0, fill=ink)
        canvas.create_line(X(0), Y(FORM_TOTAL_Y + ROW_H), X(CONTENT_W), Y(FORM_TOTAL_Y + ROW_H), fill=ink)
        canvas.create_text(col_x(_colf("taux")[1]) - 3 * scale, Y(FORM_TOTAL_Y + ROW_H / 2),
                           text="TOTAL", anchor="e", font=F(3.4, bold=True), fill=ink)
        canvas.create_text(col_x(_colf("gain")[1]) - 2.5 * scale, Y(FORM_TOTAL_Y + ROW_H / 2),
                           text=fmt_montant(res.total_gain), anchor="e", font=F(3.4, bold=True), fill=cc)
        canvas.create_text(col_x(_colf("retenue")[1]) - 2.5 * scale, Y(FORM_TOTAL_Y + ROW_H / 2),
                           text=fmt_montant(res.total_retenue), anchor="e", font=F(3.4, bold=True), fill=cc)

        # خانة NET À PAYER
        ny = FORM_NET_Y
        canvas.create_text(col_x(0.60), Y(ny + 5), text="NET À PAYER", anchor="e",
                           font=F(4.2, bold=True), fill=ink)
        canvas.create_rectangle(col_x(0.62), Y(ny), X(CONTENT_W), Y(ny + 10), fill="#111111", outline="")
        canvas.create_text(X(CONTENT_W - 3), Y(ny + 5), text=fmt_montant(res.net_a_payer),
                           anchor="e", font=F(4.8, bold=True), fill="white")


def _bulletin_rows(pin, res):
    """أسطر الجدول بالترتيب المرجعي — قائمة صفوف مفاتيحها = مفاتيح COLS."""
    c = PAIE_DEFAULT_CODES
    rows = []
    rows.append({
        "code": c["salaire_base"], "libelle": "SALAIRE DE BASE",
        "nbase": _n(pin.jours), "taux": "", "gain": fmt_montant(pin.salaire_base),
        "retenue": "",
    })
    for p in pin.primes:
        if not (p.libelle.strip() or p.montant):
            continue
        rows.append({
            "code": (p.code or c["prime"]).strip(),
            "libelle": p.libelle.strip().upper() or "PRIME",
            "nbase": "", "taux": "",
            "gain": fmt_montant(p.montant), "retenue": "",
        })
    rows.append({
        "code": c["cnas"], "libelle": "RETENUE SÉCU. SOCIALE",
        "nbase": fmt_montant(res.base_cnas), "taux": "9,00",
        "gain": "", "retenue": fmt_montant(res.retenue_cnas),
    })
    rows.append({
        "code": c["irg"], "libelle": "RETENUE IRG",
        "nbase": fmt_montant(res.base_irg), "taux": "",
        "gain": "", "retenue": fmt_montant(res.retenue_irg),
    })
    if pin.panier:
        rows.append({
            "code": c["panier"], "libelle": "PANIER", "nbase": "", "taux": "",
            "gain": fmt_montant(pin.panier), "retenue": "",
        })
    if pin.transport:
        rows.append({
            "code": c["transport"], "libelle": "(R+) TRANSPORT", "nbase": "", "taux": "",
            "gain": fmt_montant(pin.transport), "retenue": "",
        })
    for r in pin.autres_retenues:
        if not (r.libelle.strip() or r.montant):
            continue
        rows.append({
            "code": r.code.strip(), "libelle": r.libelle.strip().upper() or "RETENUE",
            "nbase": "", "taux": "", "gain": "", "retenue": fmt_montant(r.montant),
        })
    return rows


def _n(x):
    try:
        f = float(x)
    except (TypeError, ValueError):
        return ""
    return str(int(f)) if f == int(f) else f"{f:.2f}"


# ============================================================================
#  مصدر الحقيقة الموحّد للمُصيِّر (Phase C2): BulletinView — نفس محرّك الشاشة
#  (‏lignes.compute_bulletin). لا حساب هنا؛ المجاميع/الصافي من view.result.
#
#  Absence/Retard (Z1 — RETENUE خاضعة) تُنقِص TOTAL_GAINS [6] في المحرّك،
#  ولا تدخل TOTAL_RETENUES [11]. فالتمثيل المتّسق حسابياً: **مبلغ سالب في
#  عمود GAIN** — عندها Σعمود GAIN − Σعمود RETENUE = NET بالضبط.
# ============================================================================
_VIEW_ORDER = {
    "salaire_base": 0, "iep": 10, "pri": 12, "nuit": 15,
    "hs_50": 30, "hs_100": 31, "conge_paye": 35,
    "abs_jours": 40, "abs_heures": 41, "retard": 50,
    "panier": 60, "transport": 70, "alloc_fam": 75,
    "avance": 100, "syndicat": 105,
}
_BASE_REDUCERS = {"abs_jours", "abs_heures", "retard"}

#  المستند فرنسيّ — تسميات أنواع المحرّك عربية في LINE_TYPES، فنُترجمها
#  هنا (نفس تسميات شاشة ui2). السطر الحرّ (libre) يحتفظ بتسمية المستخدم.
_FR_LIBELLE = {
    "salaire_base": "SALAIRE DE BASE",
    "iep": "IND. EXPÉRIENCE PROF.", "pri": "PRIME DE RENDEMENT",
    "hs_50": "HEURES SUPP. 50 %", "hs_100": "HEURES SUPP. 100 %",
    "abs_jours": "ABSENCE (JOURS)", "abs_heures": "ABSENCE (HEURES)",
    "retard": "RETARD", "panier": "PANIER", "transport": "(R+) TRANSPORT",
    "avance": "AVANCE / ACOMPTE", "syndicat": "COTISATION SYNDICALE",
    "nuit": "PRIME DE NUIT", "conge_paye": "CONGÉ PAYÉ",
    "alloc_fam": "ALLOCATIONS FAMILIALES",
}


def _view_order(lv):
    o = _VIEW_ORDER.get(lv.key)
    if o is not None:
        return o
    #  سطر حرّ (libre): مكسب ≈ Prime · اقتطاع ≈ Avance
    return 20 if lv.sens != "RETENUE" else 100


def _bulletin_rows_from_view(view, employee, *, jours=None):
    """أسطر جدول الكشف من :class:`~programme.payroll.lignes.BulletinView`.

    كلّ رُبريكة ديناميكية في ``view.lignes`` تُمثَّل؛ CNAS/IRG سطران
    نظاميّان من ‎[A]/[B]‎ و‎[C]/[D]‎؛ Panier/Transport يظهران دائماً حتى
    بصفر (§22). لا خلايا بلا معنى."""
    c = PAIE_DEFAULT_CODES
    res = view.result
    gains, retenues = [], []
    seen_pt = set()
    for lv in sorted(view.lignes, key=_view_order):
        neg = lv.key in _BASE_REDUCERS
        lib = (_FR_LIBELLE.get(lv.key)
               or (lv.libelle or "").strip().upper() or lv.key.upper())
        row = {"code": (lv.code or "").strip(), "libelle": lib,
               "nbase": _n(jours) if lv.key == "salaire_base" else "",
               "taux": "", "gain": "", "retenue": "", "_ord": _view_order(lv)}
        amt = fmt_montant(-lv.montant if neg else lv.montant)
        if lv.key in ("panier", "transport"):
            seen_pt.add(lv.key)
            row["gain"] = amt
            gains.append(row)
        elif neg or lv.sens != "RETENUE":
            row["gain"] = amt
            gains.append(row)
        else:
            row["retenue"] = amt
            retenues.append(row)
    for key, code, lab, ordv in (("panier", c["panier"], "PANIER", 60),
                                 ("transport", c["transport"],
                                  "(R+) TRANSPORT", 70)):
        if key not in seen_pt:
            gains.append({"code": code, "libelle": lab, "nbase": "", "taux": "",
                          "gain": fmt_montant(0), "retenue": "", "_ord": ordv})
    gains.sort(key=lambda r: r["_ord"])
    rows = [{k: v for k, v in r.items() if k != "_ord"} for r in gains]
    rows.append({"code": c["cnas"], "libelle": "RETENUE SÉCU. SOCIALE",
                 "nbase": fmt_montant(res.assiette_cnas), "taux": "9,00",
                 "gain": "", "retenue": fmt_montant(res.retenue_cnas)})
    rows.append({"code": c["irg"], "libelle": "RETENUE IRG",
                 "nbase": fmt_montant(res.assiette_irg), "taux": "",
                 "gain": "", "retenue": fmt_montant(res.irg)})
    rows += [{k: v for k, v in r.items() if k != "_ord"} for r in retenues]
    return rows


def _period_label(pin):
    m = pin.mois.strip()
    y = str(pin.annee).strip()
    return f"{m} {y}".strip().upper()


def _ident_pairs(employee):
    g = employee.get
    #  §23: الحالة العائلية فارغةً تُطبَع «/» **عند التصيير فقط** — القيمة
    #  الداخلية في Work Data تبقى "" (لا تُحوَّل).
    sit = (g("situation_familiale", "") or "").strip() or "/"
    return [
        ("NOM", g("nom", "")),
        ("PRÉNOM", g("prenom", "")),
        ("DATE DE NAISSANCE", g("date_naissance", "")),
        ("LIEU DE NAISSANCE", g("lieu_naissance", "")),
        ("N° SS", g("num_ss", "")),
        ("SIT. FAMILIALE", sit),
        ("DATE D'ENTRÉE", g("date_embauche", "")),
        ("MATRICULE", g("matricule", "")),
        ("FONCTION", g("fonction", "")),
    ]


class SimpleBulletinTemplate:
    # المفتاح والتسمية مصدرهما سجلّ الموديلات في طبقة الحساب (لا تكرار)
    KEY = "simple"
    LABEL = registry.get_template("simple").label

    # ================= المعاينة الحيّة (tkinter canvas) =================
    @staticmethod
    def paint(canvas, page_box, scale, pin, res, employer, employee):
        x0, y0, _x1, _y1 = page_box

        def X(mm):
            return x0 + (MARGIN_L + mm) * scale

        def Y(mm):
            return y0 + mm * scale

        def F(mm_h, bold=False):
            size = max(int(mm_h * scale * 0.62), 6)
            return (FORM_FONT, size, "bold") if bold else (FORM_FONT, size)

        def col_x(frac):
            return x0 + (MARGIN_L + frac * CONTENT_W) * scale

        ink = "#202124"

        # --- رأس المكتب (بيانات المستخدم، بلا شعار) ---
        eg = employer.get
        canvas.create_text(X(0), Y(12), text=eg("raison_sociale", "") or "—",
                           anchor="w", font=F(5.2, bold=True), fill=ink)
        canvas.create_text(X(0), Y(20), text=eg("adresse", ""), anchor="w",
                           font=F(3.4), fill=ink)
        adh = eg("cnas_employeur", "")
        canvas.create_text(X(0), Y(26), text=(f"N° ADHÉRENT   {adh}" if adh else "N° ADHÉRENT"),
                           anchor="w", font=F(3.4), fill=ink)

        # --- شريط العنوان الأسود ---
        canvas.create_rectangle(X(0), Y(BAND_TITLE_Y), X(CONTENT_W), Y(BAND_TITLE_Y + BAND_TITLE_H),
                                fill="#111111", outline="")
        canvas.create_text(X(3), Y(BAND_TITLE_Y + BAND_TITLE_H / 2), text="BULLETIN DE PAIE",
                           anchor="w", font=F(5.0, bold=True), fill="white")
        canvas.create_text(X(CONTENT_W / 2 + 10), Y(BAND_TITLE_Y + BAND_TITLE_H / 2),
                           text=_period_label(pin), anchor="w", font=F(4.6, bold=True), fill="white")

        # --- صندوق هوية الأجير ---
        canvas.create_rectangle(X(0), Y(IDENT_Y), X(CONTENT_W), Y(IDENT_Y + IDENT_H),
                                outline=ink)
        pairs = _ident_pairs(employee)
        for i, (lab, val) in enumerate(pairs):
            row = i // 2
            colmm = 4 + (i % 2) * (CONTENT_W / 2)
            yy = IDENT_Y + 9 + row * 11
            canvas.create_text(X(colmm), Y(yy), text=f"{lab} :", anchor="w",
                               font=F(3.1, bold=True), fill=ink)
            canvas.create_text(X(colmm + 34), Y(yy), text=val, anchor="w",
                               font=F(3.3), fill=ink)
        if employee.get("handicape_retraite") or getattr(pin, "handicape_retraite", False):
            canvas.create_text(X(CONTENT_W - 2), Y(IDENT_Y + IDENT_H - 4),
                               text="☑ Handicapé / Retraité", anchor="e", font=F(3.0), fill=ink)

        # --- ترويسة الجدول ---
        hy0, hy1 = Y(TABLE_HEAD_Y), Y(TABLE_HEAD_Y + ROW_H + 1)
        canvas.create_rectangle(X(0), hy0, X(CONTENT_W), hy1, outline=ink)
        for _key, f0, _f1, _al, title in COLS:
            canvas.create_line(col_x(f0), hy0, col_x(f0), Y(TABLE_BOTTOM), fill=ink)
            canvas.create_text(col_x(f0) + 3 * scale, (hy0 + hy1) / 2, text=title,
                               anchor="w", font=F(3.1, bold=True), fill=ink)
        canvas.create_line(X(CONTENT_W), hy0, X(CONTENT_W), Y(TABLE_BOTTOM), fill=ink)

        # --- أسطر الجدول ---
        rows = _bulletin_rows(pin, res)
        ry = TABLE_HEAD_Y + ROW_H + 1
        for r in rows:
            yc = Y(ry + ROW_H / 2)
            for key, f0, f1, al, _title in COLS:
                txt = r.get(key, "")
                if not txt:
                    continue
                if al == "r":
                    canvas.create_text(col_x(f1) - 3 * scale, yc, text=txt, anchor="e",
                                       font=F(3.2), fill=ink)
                elif al == "c":
                    canvas.create_text((col_x(f0) + col_x(f1)) / 2, yc, text=txt, anchor="center",
                                       font=F(3.2), fill=ink)
                else:
                    canvas.create_text(col_x(f0) + 3 * scale, yc, text=txt, anchor="w",
                                       font=F(3.2), fill=ink)
            ry += ROW_H

        # --- سطر TOTAL ---
        ty = max(ry + 2, TABLE_BOTTOM - 2 * ROW_H)
        canvas.create_line(X(0), Y(TABLE_BOTTOM), X(CONTENT_W), Y(TABLE_BOTTOM), fill=ink)
        canvas.create_line(X(0), Y(ty), X(CONTENT_W), Y(ty), fill=ink)
        canvas.create_text(col_x(COLS[3][2]) - 3 * scale, Y(ty + ROW_H / 2), text="TOTAL",
                           anchor="e", font=F(3.4, bold=True), fill=ink)
        canvas.create_text(col_x(COLS[4][2]) - 3 * scale, Y(ty + ROW_H / 2),
                           text=fmt_montant(res.total_gain), anchor="e", font=F(3.4, bold=True), fill=ink)
        canvas.create_text(col_x(COLS[5][2]) - 3 * scale, Y(ty + ROW_H / 2),
                           text=fmt_montant(res.total_retenue), anchor="e", font=F(3.4, bold=True), fill=ink)

        # --- خانة NET À PAYER السوداء ---
        ny0 = TABLE_BOTTOM + 6
        canvas.create_text(col_x(0.60), Y(ny0 + 5), text="NET À PAYER", anchor="e",
                           font=F(4.4, bold=True), fill=ink)
        canvas.create_rectangle(col_x(0.62), Y(ny0), X(CONTENT_W), Y(ny0 + 10), fill="#111111", outline="")
        canvas.create_text(X(CONTENT_W - 3), Y(ny0 + 5), text=fmt_montant(res.net_a_payer),
                           anchor="e", font=F(5.0, bold=True), fill="white")

    # ===================== إخراج Word (.docx) =====================
    @staticmethod
    def build_docx(path, pin, res, employer, employee, *, view=None):
        try:
            from docx import Document
            from docx.shared import Pt, Mm, RGBColor
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            from docx.enum.table import WD_TABLE_ALIGNMENT
            from docx.oxml.ns import qn
            from docx.oxml import OxmlElement
        except ImportError as exc:
            raise TemplateNotReady("مكتبة python-docx غير مثبّتة — تعذّر توليد ملف Word.") from exc

        def shade(cell, hex_fill):
            tcpr = cell._tc.get_or_add_tcPr()
            sh = OxmlElement("w:shd")
            sh.set(qn("w:val"), "clear")
            sh.set(qn("w:fill"), hex_fill)
            tcpr.append(sh)

        doc = Document()
        sec = doc.sections[0]
        sec.page_width, sec.page_height = Mm(210), Mm(297)
        sec.left_margin = sec.right_margin = Mm(MARGIN_L)
        sec.top_margin = sec.bottom_margin = Mm(MARGIN_T)

        eg = employer.get
        h = doc.add_paragraph()
        r = h.add_run(eg("raison_sociale", "") or "—")
        r.bold = True
        r.font.size = Pt(14)
        if eg("adresse", ""):
            doc.add_paragraph(eg("adresse", "")).runs[0].font.size = Pt(9)
        adh = eg("cnas_employeur", "")
        doc.add_paragraph(f"N° ADHÉRENT   {adh}".rstrip()).runs[0].font.size = Pt(9)

        # شريط العنوان
        t = doc.add_table(rows=1, cols=2)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = True
        c0, c1 = t.rows[0].cells
        for cell, txt, align in ((c0, "BULLETIN DE PAIE", WD_ALIGN_PARAGRAPH.LEFT),
                                 (c1, _period_label(pin), WD_ALIGN_PARAGRAPH.RIGHT)):
            shade(cell, "111111")
            p = cell.paragraphs[0]
            p.alignment = align
            run = p.add_run(txt)
            run.bold = True
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

        doc.add_paragraph()

        # هوية الأجير
        pairs = _ident_pairs(employee)
        it = doc.add_table(rows=(len(pairs) + 1) // 2, cols=4)
        it.style = "Table Grid"
        for i, (lab, val) in enumerate(pairs):
            cell = it.rows[i // 2].cells[(i % 2) * 2]
            cell.text = lab
            cell.paragraphs[0].runs[0].bold = True
            it.rows[i // 2].cells[(i % 2) * 2 + 1].text = val or ""
        if employee.get("handicape_retraite") or getattr(pin, "handicape_retraite", False):
            doc.add_paragraph("☑ Handicapé / Retraité").runs[0].font.size = Pt(9)

        doc.add_paragraph()

        # جدول الرُّبريكات — من BulletinView إن مُرّر (مصدر الحقيقة، Phase C2)
        rows = (_bulletin_rows_from_view(view, employee, jours=pin.jours)
                if view is not None else _bulletin_rows(pin, res))
        rt = doc.add_table(rows=1 + len(rows) + 2, cols=6)
        rt.style = "Table Grid"
        heads = [c[4] for c in COLS]
        for j, htxt in enumerate(heads):
            cell = rt.rows[0].cells[j]
            cell.text = htxt
            cell.paragraphs[0].runs[0].bold = True
        for i, rowd in enumerate(rows, start=1):
            for j, (key, _f0, _f1, al, _t) in enumerate(COLS):
                cell = rt.rows[i].cells[j]
                cell.text = rowd.get(key, "")
                if al == "r":
                    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
                elif al == "c":
                    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

        # سطر TOTAL
        tr = rt.rows[1 + len(rows)]
        merged = tr.cells[0].merge(tr.cells[3])
        merged.text = "TOTAL"
        merged.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        merged.paragraphs[0].runs[0].bold = True
        for j, val in ((4, res.total_gain), (5, res.total_retenue)):
            tr.cells[j].text = fmt_montant(val)
            tr.cells[j].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
            tr.cells[j].paragraphs[0].runs[0].bold = True

        # سطر NET À PAYER
        nr = rt.rows[2 + len(rows)]
        nm = nr.cells[0].merge(nr.cells[4])
        nm.text = "NET À PAYER"
        nm.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        nm.paragraphs[0].runs[0].bold = True
        shade(nr.cells[5], "111111")
        nc = nr.cells[5].paragraphs[0]
        nc.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = nc.add_run(fmt_montant(res.net_a_payer))
        run.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

        os.makedirs(os.path.dirname(path), exist_ok=True)
        doc.save(path)
        return path

    # ===================== إخراج PDF =====================
    @staticmethod
    def build_pdf(path, pin, res, employer, employee, *, view=None):
        try:
            from reportlab.lib.pagesizes import A4
            from reportlab.lib.units import mm
            from reportlab.lib import colors
            from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
            from reportlab.platypus import (
                SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle,
            )
        except ImportError as exc:
            raise TemplateNotReady("مكتبة reportlab غير مثبّتة — تعذّر توليد ملف PDF.") from exc

        os.makedirs(os.path.dirname(path), exist_ok=True)
        styles = getSampleStyleSheet()
        small = ParagraphStyle("small", parent=styles["Normal"], fontSize=8, leading=10)
        big = ParagraphStyle("big", parent=styles["Normal"], fontSize=13, leading=15, spaceAfter=2)

        doc = SimpleDocTemplate(
            path, pagesize=A4,
            leftMargin=MARGIN_L * mm, rightMargin=MARGIN_R * mm,
            topMargin=MARGIN_T * mm, bottomMargin=MARGIN_T * mm,
        )
        eg = employer.get
        story = [
            Paragraph(f"<b>{eg('raison_sociale', '') or '—'}</b>", big),
        ]
        if eg("adresse", ""):
            story.append(Paragraph(eg("adresse", ""), small))
        adh = eg("cnas_employeur", "")
        story.append(Paragraph(f"N° ADHÉRENT&nbsp;&nbsp;&nbsp;{adh}".rstrip(), small))
        story.append(Spacer(1, 6 * mm))

        full_w = (210 - MARGIN_L - MARGIN_R) * mm
        title = Table([["BULLETIN DE PAIE", _period_label(pin)]],
                      colWidths=[full_w * 0.5, full_w * 0.5])
        title.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#111111")),
            ("TEXTCOLOR", (0, 0), (-1, -1), colors.white),
            ("FONTNAME", (0, 0), (-1, -1), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 11),
            ("ALIGN", (1, 0), (1, 0), "RIGHT"),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ]))
        story.append(title)
        story.append(Spacer(1, 4 * mm))

        pairs = _ident_pairs(employee)
        id_data = []
        for i in range(0, len(pairs), 2):
            left = pairs[i]
            right = pairs[i + 1] if i + 1 < len(pairs) else ("", "")
            id_data.append([f"{left[0]} :", left[1], f"{right[0]} :" if right[0] else "", right[1]])
        idt = Table(id_data, colWidths=[full_w * 0.18, full_w * 0.32, full_w * 0.18, full_w * 0.32])
        idt.setStyle(TableStyle([
            ("BOX", (0, 0), (-1, -1), 0.6, colors.black),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
            ("FONTNAME", (2, 0), (2, -1), "Helvetica-Bold"),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]))
        story.append(idt)
        if employee.get("handicape_retraite") or getattr(pin, "handicape_retraite", False):
            story.append(Paragraph("☑ Handicapé / Retraité", small))
        story.append(Spacer(1, 4 * mm))

        rows = (_bulletin_rows_from_view(view, employee, jours=pin.jours)
                if view is not None else _bulletin_rows(pin, res))
        data = [[c[4] for c in COLS]]
        for r in rows:
            data.append([r.get(k, "") for k, *_ in COLS])
        data.append(["TOTAL", "", "", "", fmt_montant(res.total_gain), fmt_montant(res.total_retenue)])
        data.append(["NET À PAYER", "", "", "", "", fmt_montant(res.net_a_payer)])

        cw = [full_w * (c[2] - c[1]) for c in COLS]
        rt = Table(data, colWidths=cw)
        n = len(rows)
        rt.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, n), 0.5, colors.black),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("ALIGN", (2, 1), (5, n), "RIGHT"),
            ("ALIGN", (0, 1), (0, n), "CENTER"),
            # TOTAL
            ("SPAN", (0, n + 1), (3, n + 1)),
            ("ALIGN", (0, n + 1), (3, n + 1), "RIGHT"),
            ("ALIGN", (4, n + 1), (5, n + 1), "RIGHT"),
            ("FONTNAME", (0, n + 1), (-1, n + 1), "Helvetica-Bold"),
            ("LINEABOVE", (0, n + 1), (-1, n + 1), 0.8, colors.black),
            ("LINEBELOW", (4, n + 1), (5, n + 1), 0.8, colors.black),
            # NET À PAYER
            ("SPAN", (0, n + 2), (4, n + 2)),
            ("ALIGN", (0, n + 2), (4, n + 2), "RIGHT"),
            ("ALIGN", (5, n + 2), (5, n + 2), "RIGHT"),
            ("FONTNAME", (0, n + 2), (-1, n + 2), "Helvetica-Bold"),
            ("FONTSIZE", (5, n + 2), (5, n + 2), 11),
            ("BACKGROUND", (5, n + 2), (5, n + 2), colors.HexColor("#111111")),
            ("TEXTCOLOR", (5, n + 2), (5, n + 2), colors.white),
            ("TOPPADDING", (0, n + 2), (-1, n + 2), 6),
            ("BOTTOMPADDING", (0, n + 2), (-1, n + 2), 6),
        ]))
        story.append(rt)
        doc.build(story)
        return path


# --- ربط مفتاح الموديل بمُصيّره (طبقة الواجهة فقط) ---
# طبقة الحساب (programme.payroll.registry) تعرف الموديلات المتاحة
# (مفتاح + تسمية) بلا أي اعتماد على tkinter؛ هنا نربط كل مفتاح بالصنف
# الذي يرسمه ويولّد Word/PDF منه.
RENDERERS = {
    SimpleBulletinTemplate.KEY: SimpleBulletinTemplate,
}


def get_renderer(key=None):
    """صنف المُصيّر للمفتاح المطلوب (أو الافتراضي)."""
    return RENDERERS.get(key or SimpleBulletinTemplate.KEY, SimpleBulletinTemplate)
